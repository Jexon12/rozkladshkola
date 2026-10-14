import json,re
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'Домашні завдання'
OUT.mkdir(exist_ok=True)
records=json.loads((ROOT/'homework_extracted.json').read_text(encoding='utf-8'))
for r in records:
    r['homework']=re.split(r'\n\s*(?:ДОДАТКОВ|ДОДАТОК|ДОДАТКИ)',r['homework'],flags=re.I)[0].strip()
def link(p,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'466078');pr.append(col)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'16');pr.append(size);run.append(pr)
    t=OxmlElement('w:t');t.text='Конспект уроку';run.append(t);h.append(run);p._p.append(h)
for cls in [7,8,9]:
    doc=Document();s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.line_spacing=1.06
    for name in ['Title','Heading 1','Heading 2']:
        doc.styles[name].font.name='Calibri';doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Title'].font.size=Pt(23)
    doc.styles['Heading 1'].font.size=Pt(15)
    doc.styles['Heading 2'].font.size=Pt(11)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(7)
    doc.styles['Heading 2'].paragraph_format.space_after=Pt(3)
    doc.add_paragraph(f'Домашні завдання з фізики {cls} клас НУШ','Title')
    doc.add_paragraph('Завдання за уроками' if cls!=9 else 'Завдання до уроків 1–48')
    selected=sorted([r for r in records if r['cls']==cls],key=lambda r:int(re.search(r'Урок\s+(\d+)',r['lesson']).group(1)))
    semester=None
    for r in selected:
        current=r['semester']
        if current and current!=semester:
            p=doc.add_paragraph('Перший семестр' if current in ['1 семестр','I семестр'] else 'Другий семестр','Heading 1')
            if semester is not None:p.paragraph_format.page_break_before=True
            semester=current
        doc.add_paragraph(r['lesson'],'Heading 2')
        p=doc.add_paragraph(r['homework'] or '—');p.paragraph_format.keep_with_next=True
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);link(p,r['url'])
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    path=OUT/f'Домашні завдання {cls} клас НУШ.docx';doc.save(path)
    assert len(selected)=={7:69,8:70,9:47}[cls]
    readback=Document(path)
    assert len([p for p in readback.paragraphs if p.style.name=='Heading 2'])==len(selected)
    for r in selected:
        if r['homework']: assert r['homework'] in [p.text for p in readback.paragraphs]
    print(f'{cls}: {len(selected)} уроків; {path}')
(ROOT/'homework_extracted.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
